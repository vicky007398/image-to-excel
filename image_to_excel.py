"""
Picture -> Excel table converter (works fully OFFLINE).

Reads a photo / screenshot of a table (e.g. a recipe table), finds the text with an
offline OCR engine, rebuilds the rows and columns (using the table's grid lines when
they are visible), and saves an .xlsx file with the same layout.

Usage:
    python image_to_excel.py                 -> opens a small window (pick pictures)
    python image_to_excel.py pic1.jpg pic2.png   -> converts them directly
"""
import os
import sys
import re
import statistics

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ocr_engine = None


def get_ocr():
    """Load the offline OCR engine once (models are bundled inside the package)."""
    global _ocr_engine
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    return _ocr_engine


# --------------------------------------------------------------------------- OCR
def read_image(path):
    # imdecode (instead of imread) so paths with spaces / non-English letters work on Windows
    data = np.fromfile(path, dtype=np.uint8)
    img = cv2.imdecode(data, cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError(f"Cannot open image: {path}")
    # small pictures give poor OCR -> enlarge
    h, w = img.shape[:2]
    if max(h, w) < 1400:
        s = 1400 / max(h, w)
        img = cv2.resize(img, None, fx=s, fy=s, interpolation=cv2.INTER_CUBIC)
    return img


def run_ocr(img):
    result, _ = get_ocr()(img)
    items = []
    for box, text, score in (result or []):
        xs = [p[0] for p in box]
        ys = [p[1] for p in box]
        items.append({
            "text": str(text).strip(),
            "x0": min(xs), "x1": max(xs),
            "y0": min(ys), "y1": max(ys),
            "cx": (min(xs) + max(xs)) / 2,
            "cy": (min(ys) + max(ys)) / 2,
        })
    return [i for i in items if i["text"]]


# ------------------------------------------------------------- grid line detection
def _cluster_positions(idx, gap=6):
    """Turn a list of pixel indexes into centre positions of line groups."""
    if len(idx) == 0:
        return []
    groups, cur = [], [idx[0]]
    for v in idx[1:]:
        if v - cur[-1] <= gap:
            cur.append(v)
        else:
            groups.append(cur)
            cur = [v]
    groups.append(cur)
    return [int(np.mean(g)) for g in groups]


def detect_grid(img):
    """Return (row_lines_y, col_lines_x) if the picture has visible table borders."""
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    bw = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C,
                               cv2.THRESH_BINARY_INV, 15, 10)
    h, w = bw.shape
    hk = cv2.getStructuringElement(cv2.MORPH_RECT, (max(w // 25, 20), 1))
    vk = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(h // 25, 20)))
    hor = cv2.morphologyEx(bw, cv2.MORPH_OPEN, hk)
    ver = cv2.morphologyEx(bw, cv2.MORPH_OPEN, vk)

    rows = np.where(hor.sum(axis=1) > 255 * w * 0.35)[0]
    cols = np.where(ver.sum(axis=0) > 255 * h * 0.35)[0]
    ys = _cluster_positions(rows)
    xs = _cluster_positions(cols)
    if len(ys) >= 2 and len(xs) >= 2:
        return ys, xs
    return None


def build_from_grid(items, ys, xs):
    nrows, ncols = len(ys) - 1, len(xs) - 1
    grid = [[[] for _ in range(ncols)] for _ in range(nrows)]
    for it in items:
        r = next((i for i in range(nrows) if ys[i] <= it["cy"] < ys[i + 1]), None)
        c = next((j for j in range(ncols) if xs[j] <= it["cx"] < xs[j + 1]), None)
        if r is None or c is None:
            continue
        grid[r][c].append(it)
    table = []
    for row in grid:
        out = []
        for cell in row:
            cell.sort(key=lambda t: (round(t["cy"] / 8), t["x0"]))
            out.append(" ".join(t["text"] for t in cell))
        table.append(out)
    return table


# ---------------------------------------------------- no grid lines: use text layout
def build_from_layout(items):
    if not items:
        return []
    med_h = statistics.median(i["y1"] - i["y0"] for i in items)

    # rows: group by vertical centre
    items = sorted(items, key=lambda t: t["cy"])
    rows, cur = [], [items[0]]
    for it in items[1:]:
        if abs(it["cy"] - statistics.mean(c["cy"] for c in cur)) <= med_h * 0.6:
            cur.append(it)
        else:
            rows.append(cur)
            cur = [it]
    rows.append(cur)

    # columns: merge overlapping horizontal spans across the whole picture
    spans = sorted((i["x0"], i["x1"]) for i in items)
    cols = [list(spans[0])]
    for a, b in spans[1:]:
        if a <= cols[-1][1] + med_h * 0.5:
            cols[-1][1] = max(cols[-1][1], b)
        else:
            cols.append([a, b])

    # a long title/heading would glue everything into one column -> use row-wise gaps
    if len(cols) < 2:
        cols = _columns_from_rows(rows, med_h)

    table = []
    for row in rows:
        row.sort(key=lambda t: t["x0"])
        line = [""] * len(cols)
        for it in row:
            j = min(range(len(cols)), key=lambda k: abs(it["cx"] - (cols[k][0] + cols[k][1]) / 2))
            line[j] = (line[j] + " " + it["text"]).strip()
        table.append(line)
    return table


def _columns_from_rows(rows, med_h):
    best = max(rows, key=len)
    return [[i["x0"], i["x1"]] for i in sorted(best, key=lambda t: t["x0"])]


# ------------------------------------------------------------------- clean + save
def to_value(s):
    """Store numbers as real numbers so Excel can calculate with them."""
    t = s.replace(",", "").strip()
    if re.fullmatch(r"-?\d+", t) and not (len(t) > 1 and t.startswith("0")):
        return int(t)
    if re.fullmatch(r"-?\d*\.\d+", t):
        return float(t)
    return s


def tidy(table):
    table = [r for r in table if any(c.strip() for c in r)]
    if not table:
        return table
    keep = [j for j in range(len(table[0])) if any(r[j].strip() for r in table)]
    return [[r[j] for j in keep] for r in table]


def save_xlsx(table, path, sheet_name="Recipe"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="D9E1F2")

    for r, row in enumerate(table, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=to_value(val))
            cell.font = Font(name="Arial", size=10, bold=(r == 1))
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal="center" if r == 1 else "left")
            if r == 1:
                cell.fill = head_fill

    for c in range(1, (len(table[0]) if table else 0) + 1):
        longest = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, len(table) + 1))
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, longest + 3), 45)
    wb.save(path)


def convert(image_path, out_path=None):
    img = read_image(image_path)
    items = run_ocr(img)
    if not items:
        raise ValueError("No text found in the picture.")
    grid = detect_grid(img)
    table = build_from_grid(items, *grid) if grid else build_from_layout(items)
    table = tidy(table)
    if not table:
        raise ValueError("Could not build a table from this picture.")
    out_path = out_path or os.path.splitext(image_path)[0] + ".xlsx"
    save_xlsx(table, out_path)
    return out_path, len(table), len(table[0]), bool(grid)


# --------------------------------------------------------------------------- GUI
def run_gui():
    import threading
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext

    root = tk.Tk()
    root.title("Recipe Picture -> Excel (offline)")
    root.geometry("560x380")

    log = scrolledtext.ScrolledText(root, height=14, state="disabled")

    def write(msg):
        log.config(state="normal")
        log.insert("end", msg + "\n")
        log.see("end")
        log.config(state="disabled")

    def work(files):
        for f in files:
            write(f"Reading: {os.path.basename(f)} ...")
            try:
                out, nr, nc, grid = convert(f)
                how = "grid lines detected" if grid else "layout guessed (no grid lines)"
                write(f"  Saved: {out}\n  {nr} rows x {nc} columns, {how}")
            except Exception as e:  # show any problem instead of crashing
                write(f"  FAILED: {e}")
        write("Done. Open the .xlsx and check the numbers once.")
        btn.config(state="normal")

    def pick():
        files = filedialog.askopenfilenames(
            title="Choose recipe pictures",
            filetypes=[("Pictures", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp")])
        if files:
            btn.config(state="disabled")
            write("Loading offline OCR engine (first time takes a few seconds)...")
            threading.Thread(target=work, args=(files,), daemon=True).start()

    tk.Label(root, text="Convert recipe table pictures to Excel", font=("Arial", 13, "bold")).pack(pady=8)
    btn = tk.Button(root, text="Choose picture(s)...", font=("Arial", 11), command=pick)
    btn.pack(pady=4)
    tk.Label(root, text="The Excel file is saved next to each picture with the same name.").pack()
    log.pack(fill="both", expand=True, padx=10, pady=8)
    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) > 1:
        for p in sys.argv[1:]:
            try:
                out, nr, nc, grid = convert(p)
                print(f"{p} -> {out}  ({nr} rows x {nc} cols, {'grid' if grid else 'layout'})")
            except Exception as e:
                print(f"{p}: FAILED - {e}")
    else:
        run_gui()
